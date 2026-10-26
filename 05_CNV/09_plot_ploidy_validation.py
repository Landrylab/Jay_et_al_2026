#!/usr/bin/env python3
"""Plot allele-fraction density, genomic copy number, and flow-cytometry source tables."""

import argparse
import json
from pathlib import Path
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from matplotlib.collections import PatchCollection
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.lines import Line2D
from matplotlib import font_manager
import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from figure_sources import read_source_table


BLUE, GREY, RED = "#0072B2", "#7A7A7A", "#B22222"
TABLES = {"samples": "ED3Samples", "densities": "ED3Densities",
          "genome": "ED3GenomeBins", "wells": "ED3FlowWells", "flow_bins": "ED3FlowBins", "composition": "ED3Composition"}


def require(frame, columns, name):
    missing = set(columns).difference(frame.columns)
    if missing or frame.empty:
        raise ValueError(f"{name}: expected nonempty table with columns {sorted(columns)}; missing {sorted(missing)}")


def numbers(frame, columns, nullable=()):
    for column in columns:
        raw = frame[column]
        absent = raw.isna() | raw.astype(str).str.strip().isin(["", "NA", "NaN"])
        result = pd.to_numeric(raw.where(~absent), errors="coerce")
        if ((~absent) & result.isna()).any() or (column not in nullable and result.isna().any()):
            raise ValueError(f"Invalid or missing numeric values in {column}")
        if np.isinf(result.to_numpy(float)).any():
            raise ValueError(f"Nonfinite values in {column}")
        frame[column] = result


def text_columns(frame, columns):
    for column in columns:
        if frame[column].isna().any() or frame[column].astype(str).str.strip().eq("").any():
            raise ValueError(f"Missing identifier in {column}")
        frame[column] = frame[column].astype(str)


def validate_data(samples, density, genome, wells, flow_bins, composition=None):
    samples, density, genome, wells, flow_bins = [x.copy() for x in (samples, density, genome, wells, flow_bins)]
    require(samples, ["Panel", "Isolate", "Assigned ploidy (N)"], "samples")
    require(density, ["Panel", "Allele fraction", "REF density", "ALT density"], "densities")
    require(genome, ["Panel", "Chromosome", "Bin start (bp)", "Bin end (bp)", "Absolute copy number", "Median minor AF"], "genome")
    text_columns(samples, ["Panel", "Isolate"])
    text_columns(density, ["Panel"])
    text_columns(genome, ["Panel", "Chromosome"])
    numbers(samples, ["Assigned ploidy (N)"])
    if samples.Panel.duplicated().any() or samples.Isolate.duplicated().any():
        raise ValueError("Sample panels and isolates must be unique")
    ploidies = samples["Assigned ploidy (N)"]
    if ((ploidies < 1) | (ploidies % 1 != 0)).any():
        raise ValueError("Assigned ploidy must be a positive integer")
    panels = set(samples.Panel)
    if set(density.Panel) != panels or set(genome.Panel) != panels:
        raise ValueError("Density and genome tables must cover exactly the sample panels")
    numbers(density, ["Allele fraction", "REF density", "ALT density"])
    if not density["Allele fraction"].between(0, 1).all() or (density[["REF density", "ALT density"]] < 0).any().any():
        raise ValueError("Density coordinates must have AF in [0,1] and nonnegative densities")
    if density.duplicated(["Panel", "Allele fraction"]).any() or density.groupby("Panel").size().min() < 2:
        raise ValueError("Each density curve needs at least two unique AF coordinates")
    numbers(genome, ["Bin start (bp)", "Bin end (bp)", "Absolute copy number", "Median minor AF"], ["Median minor AF"])
    if ((genome["Bin start (bp)"] < 0) | (genome["Bin end (bp)"] <= genome["Bin start (bp)"]) | (genome["Absolute copy number"] < 0)).any():
        raise ValueError("Invalid genomic interval or copy number")
    if (genome[["Bin start (bp)", "Bin end (bp)"]] % 1 != 0).any().any():
        raise ValueError("Genomic coordinates must be integer base positions")
    if not genome["Median minor AF"].dropna().between(0, .5).all():
        raise ValueError("Minor AF must be in [0,0.5] or missing")
    if genome.duplicated(["Panel", "Chromosome", "Bin start (bp)"]).any():
        raise ValueError("Duplicated genomic bin")
    for _, group in genome.groupby(["Panel", "Chromosome"], sort=False):
        ordered = group.sort_values("Bin start (bp)")
        if np.any(ordered["Bin start (bp)"].to_numpy()[1:] < ordered["Bin end (bp)"].to_numpy()[:-1]):
            raise ValueError("Genomic bins overlap")
    require(wells, ["sample", "well", "recorded_ploidy", "group", "cohort_sample", "control_strain", "events"], "flow wells")
    require(flow_bins, ["well", "bin_start", "bin_end", "events", "percent_of_all_events"], "flow bins")
    text_columns(wells, ["sample", "well", "recorded_ploidy", "group"])
    text_columns(flow_bins, ["well"])
    if wells.well.duplicated().any() or not set(wells.group).issubset({"Cohort isolates", "Reference controls"}):
        raise ValueError("Flow wells must be unique with cohort/reference group labels")
    if not wells.recorded_ploidy.str.fullmatch(r"[1-9][0-9]*N").all():
        raise ValueError("Recorded ploidy labels must use positive integer N values, e.g. 2N")
    wells[["cohort_sample", "control_strain"]] = wells[["cohort_sample", "control_strain"]].fillna("").astype(str)
    cohort = wells.group.eq("Cohort isolates")
    if wells.loc[cohort, "cohort_sample"].eq("").any() or wells.loc[cohort, "cohort_sample"].duplicated().any():
        raise ValueError("Each cohort recording needs one unique cohort sample")
    if wells.loc[~cohort, "control_strain"].eq("").any() or wells.loc[~cohort, "cohort_sample"].ne("").any():
        raise ValueError("Reference recordings need control identities and no cohort sample")
    numbers(wells, ["events"])
    numbers(flow_bins, ["bin_start", "bin_end", "events", "percent_of_all_events"])
    if ((wells.events <= 0) | (wells.events % 1 != 0)).any() or ((flow_bins.events < 0) | (flow_bins.events % 1 != 0)).any():
        raise ValueError("Recorded events must be positive integer totals and nonnegative integer bin counts")
    if set(wells.well) != set(flow_bins.well) or flow_bins.duplicated(["well", "bin_start"]).any():
        raise ValueError("Flow-bin wells must match the well table exactly without duplicate bins")
    flow_bins = flow_bins.sort_values(["well", "bin_start"], kind="stable").reset_index(drop=True)
    common_bins = None
    totals = wells.set_index("well").events
    for well, part in flow_bins.groupby("well", sort=False):
        lo, hi = part.bin_start.to_numpy(), part.bin_end.to_numpy()
        if lo[0] < 0 or np.any(hi <= lo) or np.any(lo[1:] != hi[:-1]):
            raise ValueError(f"Flow bins for {well} must be contiguous ordered intervals")
        bins = np.column_stack([lo, hi])
        if common_bins is not None and not np.array_equal(bins, common_bins):
            raise ValueError("All flow recordings must use the same bins")
        common_bins = bins
        if int(part.events.sum()) != totals[well] or not np.allclose(part.percent_of_all_events, 100 * part.events / totals[well], rtol=1e-12, atol=1e-12):
            raise ValueError(f"Flow percentages for {well} must include all recorded events")
    wells["_ploidy"] = wells.recorded_ploidy.str[:-1].astype(int)
    wells["_group"] = wells.group.map({"Cohort isolates": 0, "Reference controls": 1})
    wells = wells.sort_values(["_ploidy", "_group"], kind="stable").reset_index(drop=True)
    mapped = []
    for _, item in samples.iterrows():
        if "Flow well" in samples:
            match = "" if pd.isna(item["Flow well"]) else str(item["Flow well"])
            if match and match not in set(wells.well):
                raise ValueError(f"Mapped flow well {match!r} is absent")
        else:
            matches = wells.loc[wells.cohort_sample.eq(item.Isolate), "well"].tolist()
            match = matches[0] if matches else ""
        mapped.append(match)
    if len([well for well in mapped if well]) != len(set(well for well in mapped if well)):
        raise ValueError("Each example must map to a distinct flow well")
    samples["Flow well"] = mapped
    if composition is not None:
        composition = composition.copy()
        require(composition, ["Chromosome", "Aneuploidy type", "Copy number", "Isolates", "Chromosome denominator", "Aneuploid isolates (%)"], "composition")
        text_columns(composition, ["Chromosome", "Aneuploidy type"])
        numbers(composition, ["Copy number", "Isolates", "Chromosome denominator", "Aneuploid isolates (%)"])
        if not set(composition["Aneuploidy type"]).issubset({"whole chromosome", "segmental"}):
            raise ValueError("Unknown aneuploidy type")
        columns = ["Copy number", "Isolates", "Chromosome denominator"]
        if (composition[columns] < 0).any().any() or (composition[columns] % 1 != 0).any().any():
            raise ValueError("Composition counts and copy numbers must be nonnegative integers")
        if composition.duplicated(["Chromosome", "Aneuploidy type", "Copy number"]).any():
            raise ValueError("Duplicated composition category")
        for chrom, part in composition.groupby("Chromosome", sort=False):
            denominator = part["Chromosome denominator"].iloc[0]
            if denominator <= 0 or part["Chromosome denominator"].nunique() != 1 or part.Isolates.sum() != denominator:
                raise ValueError(f"Composition counts do not reconcile for chromosome {chrom}")
            if not np.allclose(part["Aneuploid isolates (%)"], 100 * part.Isolates / denominator, rtol=1e-12, atol=1e-12):
                raise ValueError(f"Incorrect composition percentages for chromosome {chrom}")
    return samples, density, genome, wells, flow_bins, composition


def load_table(source, name):
    """Read named CSV/XLSX tables or sections in the current submission workbook."""
    sections = {"ED3FlowWells": "Panel e: flow cytometry wells in heatmap order",
                "ED3FlowBins": "Panel e: fluorescence bins",
                "ED3Composition": "Panel f: chromosome-specific aneuploidy composition"}
    try:
        return read_source_table(source, name)
    except ValueError as exc:
        if Path(source).suffix.lower() != ".xlsx" or name not in sections or not str(exc).startswith("Missing Excel table"):
            raise
    from openpyxl import load_workbook
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = "Flow cytometry and composition"
        if sheet not in workbook.sheetnames:
            raise ValueError(f"Expected the current '{sheet}' worksheet")
        rows = list(workbook[sheet].values)
        locations = [i for i, row in enumerate(rows) if row[0] == sections[name]]
        if len(locations) != 1:
            raise ValueError(f"Expected one section: {sections[name]}")
        start = locations[0] + 1
        headers = list(rows[start])
        while headers and headers[-1] is None:
            headers.pop()
        if not headers or None in headers or len(headers) != len(set(headers)):
            raise ValueError(f"Invalid header for {name}")
        values = []
        for row in rows[start + 1:]:
            if row[0] is None:
                break
            values.append(row[:len(headers)])
        return pd.DataFrame(values, columns=headers)
    finally:
        workbook.close()


def chromosome_coordinates(genome):
    lengths = genome.groupby("Chromosome", sort=False)["Bin end (bp)"].max()
    return lengths, lengths.cumsum().shift(fill_value=0)


def plot(samples, density, genome, wells, flow_bins, composition, output_prefix,
         dpi=600, flow_min=0, flow_max=35000, colour_max=45, cn_max=8,
         chromosome_layout="per-panel", font_family="DejaVu Sans"):
    if flow_min >= flow_max:
        raise ValueError("The fluorescence minimum must be below its maximum")
    if flow_bins.percent_of_all_events.max() > colour_max + 1e-12:
        raise ValueError("Colour maximum would clip an observed bin; increase --colour-max")
    plt.rcParams.update({"font.family": font_family, "font.size": 6, "axes.labelweight": "bold",
                         "axes.linewidth": .55, "pdf.fonttype": 42, "svg.fonttype": "none"})
    orange = "#D55E00"
    n = len(samples)
    if n > 24:
        raise ValueError("At most 24 genomic examples fit the alphabetic panel scheme")
    width, panel_bottom = 215., 17 + (n - 1) * 58 + 46
    height = panel_bottom + (72 if composition is not None else 8)
    fig = plt.figure(figsize=(width / 25.4, height / 25.4))

    def label(x, y, value, size=5.5, bold=False, ha="left", colour="#111111"):
        return fig.text(x / width, 1 - y / height, str(value), fontsize=size,
                        weight="bold" if bold else "normal", ha=ha, va="center", color=colour)

    def line(x0, y0, x1, y1, colour="#111111", lw=.45):
        fig.add_artist(Line2D([x0/width, x1/width], [1-y0/height, 1-y1/height],
                             transform=fig.transFigure, color=colour, lw=lw))

    def rectangle(x, top, w, h, colour, hatch=None, edge="none"):
        fig.add_artist(Rectangle((x/width, 1-(top+h)/height), w/width, h/height,
                                transform=fig.transFigure, facecolor=colour,
                                edgecolor=edge, linewidth=.35, hatch=hatch))

    def axis(x, top, w, h):
        ax = fig.add_axes([x/width, 1-(top+h)/height, w/width, h/height])
        ax.spines[["top", "right"]].set_visible(False)
        ax.tick_params(labelsize=5, length=1.8, width=.45, pad=1.5)
        ax.xaxis.label.set_size(6)
        ax.yaxis.label.set_size(6)
        return ax

    palette = LinearSegmentedColormap.from_list("flow", [(0,"#ffffff"),(.15,"#c6dbef"),
                 (.35,"#4292c6"),(.65,"#08519c"),(1,"#03132b")], N=2048)
    colour = lambda pc: palette(np.log1p(pc) / np.log1p(colour_max))
    heat_bottom = panel_bottom - 25.5
    groups = wells.recorded_ploidy.nunique()
    row_gap = .6
    row_step = (heat_bottom - 17 - 4*groups - 3*(groups-1) + row_gap) / len(wells)
    row_height = row_step-row_gap
    if row_height < 1:
        raise ValueError("Too many heatmap rows for these genomic panels")
    y, prior_group = 17., None
    row_centres, row_text, heat_patches, heat_colours = {}, {}, [], []
    example_wells = set(samples["Flow well"]) - {""}
    for _, record in wells.iterrows():
        if record.recorded_ploidy != prior_group:
            if prior_group is not None:
                y += 3
            label(143, y, record.recorded_ploidy, size=5.5, bold=True)
            y += 4
            prior_group = record.recorded_ploidy
        midpoint = y + row_height/2
        row_centres[record.well] = midpoint
        part = flow_bins.loc[flow_bins.well.eq(record.well)]
        for cell in part.itertuples(index=False):
            lo, hi = max(flow_min, cell.bin_start), min(flow_max, cell.bin_end)
            if hi <= lo:
                continue
            x = 143 + 70*(lo-flow_min)/(flow_max-flow_min)
            w = 70*(hi-lo)/(flow_max-flow_min)
            heat_patches.append(Rectangle((x/width, 1-(y+row_height)/height),
                                          (w+.003)/width, row_height/height))
            heat_colours.append(colour(cell.percent_of_all_events))
        display = record.get("display_label", None)
        if pd.isna(display) or not str(display).strip():
            display = record["sample"] if record.group == "Cohort isolates" else record.control_strain
        text = f"{display}  {record.recorded_ploidy}"
        row_text[record.well] = label(141, midpoint, text, size=6.5,
            bold=record.well in example_wells, ha="right",
            colour=orange if record.group == "Reference controls" else "#111111")
        if record.well in example_wells:
            line(142.5, y, 213.5, y, orange, .6)
            line(142.5, y+row_height, 213.5, y+row_height, orange, .6)
        y += row_step
    if not np.isclose(y-row_gap, heat_bottom):
        raise AssertionError("Heatmap alignment failed")
    fig.add_artist(PatchCollection(heat_patches, facecolors=heat_colours, edgecolors="none",
                                  transform=fig.transFigure, zorder=-1))
    for value in np.linspace(flow_min, flow_max, 8):
        x = 143+70*(value-flow_min)/(flow_max-flow_min)
        line(x, heat_bottom, x, heat_bottom+.65, lw=.35)
        label(x, heat_bottom+2.25, f"{value/1000:g}", size=5, ha="center")
    label(178, heat_bottom+7, "Fluorescence (10³ a.u.)", bold=True, ha="center")
    for i in range(256):
        rectangle(143+70*i/256, heat_bottom+13, 70/256+.003, 2, palette(i/255))
    ticks = [0,1,5,15,45] if colour_max == 45 else np.expm1(np.linspace(0,np.log1p(colour_max),5))
    for pc in ticks:
        x = 143+70*np.log1p(pc)/np.log1p(colour_max)
        line(x, heat_bottom+15, x, heat_bottom+15.65, lw=.35)
        label(x, heat_bottom+17.25, f"{pc:g}", size=5, ha="center")
    label(178, heat_bottom+22, "Events per bin (%)", bold=True, ha="center")
    label(178, heat_bottom+25.5, "log(1 + %) colour scale", size=5, ha="center")
    label(115, 12, chr(97+n), size=8, bold=True)
    common_lengths, common_offsets = chromosome_coordinates(genome)
    summaries = []
    for i, (_, sample) in enumerate(samples.iterrows()):
        panel, ploidy, top = sample.Panel, int(sample["Assigned ploidy (N)"]), 17+i*58
        d = density.loc[density.Panel.eq(panel)].sort_values("Allele fraction")
        ax = axis(12, top, 25, 40)
        ax.plot(d["Allele fraction"], d["ALT density"], color=RED, lw=.65)
        ax.plot(d["Allele fraction"], d["REF density"], color="#1E4E9A", lw=.65)
        for af in np.arange(1,ploidy)/ploidy:
            ax.axvline(af, color="#777777", lw=.4, ls="--")
        ax.set(xlim=(0,1), xticks=[0,.5,1], xlabel="Allele fraction", ylabel="Density")
        ax.yaxis.set_label_coords(-.29,.5)
        display = sample.get("Display label", sample.Isolate.split("_",1)[0])
        label(1, top-5, panel, size=8, bold=True)
        label(12, top-5, f"{display} ({ploidy}N)", size=7, bold=True)
        accession = sample.Isolate.partition("_")[2]
        if accession.startswith(("SRR", "ERR", "DRR")):
            label(55, top-5, accession, size=5.5)
        if i == 0:
            line(12,8,16,8,RED,.65); label(17,8,"ALT",size=5)
            line(28,8,32,8,"#1E4E9A",.65); label(33,8,"REF",size=5)
        cn_ax, af_ax = axis(48,top,57,24), axis(48,top+27,57,13)
        g = genome.loc[genome.Panel.eq(panel)]
        lengths, offsets = (common_lengths,common_offsets) if chromosome_layout == "shared" else chromosome_coordinates(g)
        for j, (chrom,length) in enumerate(lengths.items()):
            values = g.loc[g.Chromosome.eq(chrom)]
            xs = values["Bin start (bp)"]+offsets[chrom]
            cn_ax.scatter(xs, values["Absolute copy number"], s=.45, color=BLUE if j%2 == 0 else GREY, alpha=.72, linewidths=0)
            valid = values["Median minor AF"].notna()
            af_ax.scatter(xs[valid], values.loc[valid,"Median minor AF"], s=.4, color="#1E4E9A", alpha=.72, linewidths=0)
            if j:
                for target in (cn_ax,af_ax):
                    target.axvline(offsets[chrom], ls=":", color="#333333", lw=.4)
        for target in (cn_ax,af_ax):
            target.set_xlim(-.003*lengths.sum(), 1.003*lengths.sum())
        cn_ax.set(ylim=(0,cn_max), yticks=np.linspace(0,cn_max,5), ylabel="Copy number", xticks=[])
        cn_ax.axhline(ploidy, color="#333333", lw=.45)
        cn_ax.yaxis.set_label_coords(-.12,.5)
        af_ax.set(ylim=(0,.5), ylabel="Minor AF", xlabel="Chromosome",
            xticks=offsets.to_numpy()+lengths.to_numpy()/2,
            xticklabels=[str(x).removeprefix("Chr") for x in lengths.index], yticks=[0,.25,.5])
        af_ax.yaxis.set_label_coords(-.12,.5)
        for af in np.arange(1,ploidy//2+1)/ploidy:
            af_ax.axhline(af, color="#777777", ls="--", lw=.35, zorder=0)
        summaries.append({"panel":panel,"isolate":sample.Isolate,"ploidy":ploidy,
            "density_points":len(d),"CN_bins":len(g),
            "minor_AF_bins":int(g["Median minor AF"].notna().sum()),
            "CN_bins_outside_axis":int((g["Absolute copy number"]>cn_max).sum()),
            "flow_well":sample["Flow well"]})
    # Measure labels in the selected portable font before placing arrow endpoints.
    fig.canvas.draw()
    renderer, arrows = fig.canvas.get_renderer(), []
    for i, (_,sample) in enumerate(samples.iterrows()):
        well = sample["Flow well"]
        if not well:
            continue
        end_y = row_centres[well]
        label_left = row_text[well].get_window_extent(renderer).x0/fig.bbox.width*width
        end_x, start_y = label_left-.7, 17+i*58+24
        if end_x <= 108:
            raise ValueError(f"Heatmap label {well} is too wide; shorten its display_label")
        elbow = min(116,end_x-1.5)
        points = [(105,start_y),(108,start_y),(elbow,end_y),(end_x,end_y)]
        for a,b in zip(points,points[1:]):
            line(*a,*b,orange,.65)
        line(end_x,end_y,end_x-1.4,end_y-.8,orange,.65)
        line(end_x,end_y,end_x-1.4,end_y+.8,orange,.65)
        arrows.append({"well":well,"points_mm":points,"label_gap_mm":.7})
    if composition is not None:
        label(1,panel_bottom+10,chr(98+n),size=8,bold=True)
        chroms = composition.Chromosome.drop_duplicates().tolist()
        states = sorted(composition["Copy number"].unique())
        colours = {0:"#053061",1:"#2166AC",3:"#B2182B",4:"#67001F"}
        for i,state in enumerate(states):
            colours.setdefault(state,plt.get_cmap("tab10")(i%10))
            rectangle(15+i*17,panel_bottom+8.9,3,1.8,colours[state])
            label(19+i*17,panel_bottom+10,f"CN{int(state)}")
        rectangle(91,panel_bottom+8.9,3,1.8,"#737373")
        label(95,panel_bottom+10,"Whole chromosome")
        rectangle(135,panel_bottom+8.9,3,1.8,"#737373",hatch="///",edge="white")
        label(139,panel_bottom+10,"Segmental (≥30 kb)")
        ax = axis(14,panel_bottom+21,198,38)
        for j,chrom in enumerate(chroms):
            bottom = 0
            for kind in ["whole chromosome","segmental"]:
                for state in states:
                    part = composition.loc[composition.Chromosome.eq(chrom) & composition["Aneuploidy type"].eq(kind) & composition["Copy number"].eq(state)]
                    if part.empty:
                        continue
                    item = part.iloc[0]
                    value = item["Aneuploid isolates (%)"]
                    ax.bar(j,value,bottom=bottom,width=.485,color=colours[state],edgecolor="white",lw=.5,hatch="///" if kind=="segmental" else None)
                    if value >= 4:
                        ax.text(j,bottom+value/2,str(int(item.Isolates)),color="white",ha="center",va="center",fontsize=5.5,weight="bold",bbox={"facecolor":colours[state],"edgecolor":"none","pad":.15})
                    bottom += value
            denominator = int(composition.loc[composition.Chromosome.eq(chrom),"Chromosome denominator"].iloc[0])
            ax.text(j,103,f"n={denominator}",ha="center",fontsize=5.5,clip_on=False)
        ax.set(xlim=(-.5,len(chroms)-.5),ylim=(0,100),yticks=range(0,101,20),xticks=range(len(chroms)),
            xticklabels=[str(x).removeprefix("Chr") for x in chroms],ylabel="Aneuploid isolates (%)",xlabel="Chromosome")
        ax.yaxis.set_label_coords(-.055,.5)
    output_prefix = Path(output_prefix)
    output_prefix.parent.mkdir(parents=True,exist_ok=True)
    for suffix in ("pdf","png"):
        fig.savefig(str(output_prefix)+"."+suffix,dpi=dpi,facecolor="white")
    plt.close(fig)
    report = {"samples":summaries,"flow_event_counts":dict(zip(wells.well,wells.events.astype(int))),
        "flow_row_order":wells.well.tolist(),"flow_groups":wells.group.value_counts().to_dict(),
        "flow_recorded_ploidy_counts":{f"{group}: {ploidy}":count for (group,ploidy),count in wells.groupby(["group","recorded_ploidy"]).size().items()},
        "flow_bin_count":len(flow_bins),"flow_normalization":"All recorded events, including those outside the displayed range",
        "flow_display_limits":[flow_min,flow_max],"heatmap_colour_range_percent":[0,colour_max],
        "heatmap_colour_transform":"log1p","heatmap_row_height_mm":row_height,"heatmap_panel_bottom_mm":panel_bottom,
        "heatmap_label_font_pt":6.5,"arrows":arrows,"CN_display_limits":[0,cn_max],"chromosome_layout":chromosome_layout,
        "composition_rows":0 if composition is None else len(composition),"dimensions_mm":[width,height],"dpi":dpi,"font_family":font_family}
    Path(str(output_prefix)+".summary.json").write_text(json.dumps(report,indent=2)+"\n")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source",type=Path,help="Current source workbook or directory of named TSV/CSV tables")
    for key,name in TABLES.items():
        flags = ["--flow-wells", "--wells"] if key == "wells" else ["--"+key.replace("_","-")]
        parser.add_argument(*flags,dest=key,type=Path,help=f"Separate TSV/CSV for {name}; overrides --source")
    parser.add_argument("--no-composition",action="store_true")
    parser.add_argument("--output-prefix",type=Path,required=True)
    parser.add_argument("--dpi",type=int,default=600)
    parser.add_argument("--flow-min",type=float,default=0)
    parser.add_argument("--flow-max",type=float,default=35000)
    parser.add_argument("--colour-max",type=float,default=45,help="Upper percentage on the log(1 + percentage) colour scale")
    parser.add_argument("--cn-max",type=float,default=8)
    parser.add_argument("--font-family",default="DejaVu Sans",help="Installed font family; DejaVu Sans is bundled with Matplotlib")
    parser.add_argument("--chromosome-layout",choices=["shared","per-panel"],default="per-panel")
    args = parser.parse_args()
    if args.dpi <= 0 or any(not np.isfinite(x) or x <= 0 for x in [args.flow_max,args.colour_max,args.cn_max]):
        parser.error("DPI, axis maxima and colour maximum must be positive")
    if not np.isfinite(args.flow_min) or args.flow_min < 0 or args.flow_min >= args.flow_max:
        parser.error("Fluorescence minimum must be nonnegative and below maximum")
    font_manager.findfont(font_manager.FontProperties(family=args.font_family),fallback_to_default=False)
    data = {}
    for key,name in TABLES.items():
        if key == "composition" and args.no_composition:
            data[key] = None
            continue
        source = getattr(args,key) or args.source
        if source is None:
            parser.error(f"Provide --{key.replace('_','-')} or --source")
        data[key] = load_table(source,name)
    checked = validate_data(data["samples"],data["densities"],data["genome"],data["wells"],data["flow_bins"],data["composition"])
    report = plot(*checked,args.output_prefix,dpi=args.dpi,flow_min=args.flow_min,flow_max=args.flow_max,
        colour_max=args.colour_max,cn_max=args.cn_max,chromosome_layout=args.chromosome_layout,font_family=args.font_family)
    print(json.dumps({"output_prefix":str(args.output_prefix),"flow_wells":len(checked[3]),"flow_bins":len(checked[4]),
        "total_events":sum(report["flow_event_counts"].values()),"dimensions_mm":report["dimensions_mm"]},indent=2))


if __name__ == "__main__":
    main()
