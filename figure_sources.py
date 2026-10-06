"""Read named figure tables from Excel workbooks or delimited files."""

from __future__ import annotations

from pathlib import Path
import csv
import gzip
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pandas as pd


def _excel_table_ranges(path: Path) -> dict[str, tuple[str, str]]:
    """Map Excel table names to worksheet names and cell ranges."""
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    relationship_key = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

    def relationships(archive, member, base):
        import posixpath

        root = ET.fromstring(archive.read(member))
        return {
            item.attrib["Id"]: (
                item.attrib["Target"].lstrip("/")
                if item.attrib["Target"].startswith("/")
                else posixpath.normpath(posixpath.join(base, item.attrib["Target"]))
            )
            for item in root
            if item.attrib.get("TargetMode") != "External"
        }

    tables = {}
    with ZipFile(path) as archive:
        sheets = ET.fromstring(archive.read("xl/workbook.xml"))
        workbook_links = relationships(archive, "xl/_rels/workbook.xml.rels", "xl")
        for sheet in sheets.findall("s:sheets/s:sheet", namespace):
            member = workbook_links[sheet.attrib[relationship_key]]
            folder, filename = member.rsplit("/", 1)
            links_member = f"{folder}/_rels/{filename}.rels"
            if links_member not in archive.namelist():
                continue
            links = relationships(archive, links_member, folder)
            for target in links.values():
                if "/tables/" not in target:
                    continue
                table = ET.fromstring(archive.read(target))
                name = table.attrib["displayName"]
                if name in tables:
                    raise ValueError(f"Duplicate Excel table name: {name}")
                if table.attrib.get("totalsRowCount", "0") != "0":
                    raise ValueError(f"Figure source tables must not contain a totals row: {name}")
                tables[name] = (sheet.attrib["name"], table.attrib["ref"])
    return tables


def read_source_table(path: str | Path, table_name: str) -> pd.DataFrame:
    """Read an Excel named table, a delimited file, or TABLE.tsv in a directory.

    Workbook sheet positions are independent of table names. Delimited files
    have one header row; text labels such as NA are retained as text. Empty
    fields are represented as missing values. Delimited columns retain text
    types so that identifiers with leading zeroes remain intact; callers
    convert their measurement columns to numeric values explicitly.
    """
    path = Path(path)
    if path.is_dir():
        matches = [path / (table_name + suffix) for suffix in (".tsv", ".tsv.gz", ".csv", ".csv.gz")]
        matches = [candidate for candidate in matches if candidate.is_file()]
        if len(matches) != 1:
            raise ValueError(f"Expected one delimited file for table {table_name}; found {len(matches)}")
        path = matches[0]
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries

        tables = _excel_table_ranges(path)
        if table_name not in tables:
            raise ValueError(f"Missing Excel table {table_name}; available tables: {', '.join(sorted(tables))}")
        sheet, reference = tables[table_name]
        first_column, first_row, last_column, last_row = range_boundaries(reference)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook[sheet].iter_rows(
                min_row=first_row, max_row=last_row,
                min_col=first_column, max_col=last_column, values_only=True,
            )
            headers = next(rows)
            if any(value is None or not str(value).strip() for value in headers):
                raise ValueError(f"Blank column header in table {table_name}")
            headers = [str(value) for value in headers]
            if len(headers) != len(set(headers)):
                raise ValueError(f"Duplicate column header in table {table_name}")
            return pd.DataFrame((row for row in rows if any(value is not None for value in row)), columns=headers)
        finally:
            workbook.close()
    lower_name = path.name.lower()
    if not lower_name.endswith((".tsv", ".tsv.gz", ".csv", ".csv.gz")):
        raise ValueError("Source must be an .xlsx workbook, a TSV/CSV file, or a table directory")
    separator = "\t" if lower_name.endswith((".tsv", ".tsv.gz")) else ","
    opener = gzip.open if lower_name.endswith(".gz") else open
    with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:
        headers = next(csv.reader(handle, delimiter=separator), [])
    if not headers or any(not header.strip() for header in headers) or len(headers) != len(set(headers)):
        raise ValueError("Delimited source tables require nonblank, unique column headers")
    return pd.read_csv(path, sep=separator, dtype=object, keep_default_na=False, na_values=[""], low_memory=False)
